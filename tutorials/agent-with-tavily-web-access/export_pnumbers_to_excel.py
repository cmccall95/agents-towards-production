"""
Export parsed pnumbers.com data to Excel.

Usage:
    python export_pnumbers_to_excel.py [json_file]
    
If no file specified, uses the most recent pnumbers_parsed_*.json file.
"""

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_DIR = Path(__file__).parent / "output"


def find_latest_parsed_file() -> Path:
    """Find the most recent parsed pnumbers JSON file."""
    files = list(OUTPUT_DIR.glob("pnumbers_parsed_*.json"))
    if not files:
        raise FileNotFoundError("No pnumbers_parsed_*.json files found in output/")
    return max(files, key=lambda f: f.stat().st_mtime)


def export_to_excel(json_path: Path) -> Path:
    """Export parsed JSON data to Excel."""
    print(f"📂 Loading: {json_path.name}")
    
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    records = data.get("records", [])
    print(f"📊 Records to export: {len(records)}")
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "P-Numbers"
    
    # Define columns (excluding raw_line for cleaner output)
    columns = [
        ("P#", "p_number", 6),
        ("Gr#", "group_number", 6),
        ("ISO", "iso_group", 8),
        ("Spec", "spec", 12),
        ("Type/Grade", "type_grade", 14),
        ("UNS", "uns_code", 10),
        ("ksi", "tensile_ksi", 8),
        ("MPa", "tensile_mpa", 8),
        ("Product Form", "product_form", 25),
        ("Thickness", "thickness", 15),
    ]
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write headers
    for col_idx, (header, _, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    # Write data rows
    for row_idx, record in enumerate(records, start=2):
        for col_idx, (_, field, _) in enumerate(columns, start=1):
            value = record.get(field)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
    
    # Freeze header row
    ws.freeze_panes = "A2"
    
    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(records) + 1}"
    
    # Save
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_path = OUTPUT_DIR / f"pnumbers_data_{timestamp}.xlsx"
    wb.save(excel_path)
    
    print(f"✅ Exported to: {excel_path.name}")
    return excel_path


def main():
    print("="*60)
    print("📊 PNUMBERS DATA EXCEL EXPORT")
    print("="*60 + "\n")
    
    if len(sys.argv) > 1:
        json_path = Path(sys.argv[1])
    else:
        json_path = find_latest_parsed_file()
    
    excel_path = export_to_excel(json_path)
    
    print(f"\n💾 Output: {excel_path}")
    print(f"📁 Location: {excel_path.parent}")
    
    return excel_path


if __name__ == "__main__":
    main()

