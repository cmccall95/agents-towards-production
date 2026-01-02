#!/usr/bin/env python3
"""
Export ASTM batch JSON metadata to Excel (XLSX) file.

This script reads all batch JSON files from the output directory and
consolidates them into a single Excel file for review and manual selection.
"""

import json
import glob
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("Error: openpyxl not installed. Run: pip install openpyxl")
    exit(1)


def load_batch_files(json_dir: Path) -> list:
    """Load all batch JSON files and extract results."""
    all_results = []
    batch_files = sorted(glob.glob(str(json_dir / "batch_*.json")))
    
    print(f"Found {len(batch_files)} batch files")
    
    for batch_file in batch_files:
        with open(batch_file, 'r', encoding='utf-8') as f:
            batch_data = json.load(f)
            
        batch_num = batch_data.get('batch_number', 'unknown')
        results = batch_data.get('results', [])
        
        # Add batch number to each result
        for result in results:
            result['batch_number'] = batch_num
            all_results.append(result)
        
        print(f"  Loaded batch {batch_num}: {len(results)} standards")
    
    return all_results


def create_excel(results: list, output_path: Path):
    """Create Excel file from results."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ASTM Standards"
    
    # Define columns
    columns = [
        ('batch_number', 'Batch'),
        ('code', 'Code'),
        ('designation', 'Designation'),
        ('title', 'Title'),
        ('url', 'URL'),
        ('content_length', 'Content Length'),
        ('markdown_file', 'Markdown File'),
        ('success', 'Extraction Success'),
        ('extract_grades', 'Extract Grades?'),  # Empty column for manual marking
        ('notes', 'Notes'),  # Empty column for notes
    ]
    
    # Write header row
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_idx, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    
    # Write data rows
    for row_idx, result in enumerate(results, start=2):
        for col_idx, (field, _) in enumerate(columns, start=1):
            value = result.get(field, '')
            if field in ('extract_grades', 'notes'):
                value = ''  # Empty for manual entry
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    # Adjust column widths
    column_widths = {
        'A': 8,   # Batch
        'B': 18,  # Code
        'C': 22,  # Designation
        'D': 60,  # Title
        'E': 45,  # URL
        'F': 15,  # Content Length
        'G': 30,  # Markdown File
        'H': 18,  # Success
        'I': 15,  # Extract Grades?
        'J': 30,  # Notes
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Freeze header row
    ws.freeze_panes = 'A2'
    
    # Add autofilter
    ws.auto_filter.ref = ws.dimensions
    
    # Save
    wb.save(output_path)
    print(f"\n✅ Excel file saved: {output_path}")
    print(f"   Total rows: {len(results)}")


def main():
    # Paths
    base_dir = Path(__file__).parent
    json_dir = base_dir / "output" / "astm_lists" / "full_summary" / "json"
    output_file = base_dir / "astm_standards_metadata.xlsx"
    
    print("=" * 60)
    print("ASTM Standards Batch Data Export to Excel")
    print("=" * 60)
    print(f"\nSource: {json_dir}")
    print(f"Output: {output_file}\n")
    
    # Load data
    results = load_batch_files(json_dir)
    
    if not results:
        print("No results found!")
        return
    
    # Create Excel
    create_excel(results, output_file)
    
    print(f"\n📊 Summary:")
    print(f"   Standards exported: {len(results)}")
    print(f"   Successful extractions: {sum(1 for r in results if r.get('success'))}")
    print(f"   Failed extractions: {sum(1 for r in results if not r.get('success'))}")


if __name__ == "__main__":
    main()

